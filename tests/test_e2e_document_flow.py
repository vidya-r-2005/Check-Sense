"""End-to-End SRS Document Upload, Batch Analysis, and Refinement Smoke Test.

Executes complete lifecycle verification:
1. Document ingestion (.txt and .docx)
2. Requirement segmentation, section identification, and metadata tracking
3. Batch multi-head defect detection (Heads 4a, 4b, 4c) and fusion scoring
4. Grounded question formulation
5. Stakeholder answer input
6. Governed requirement refinement with numeric preservation audit
7. Explicit stakeholder approval gating
8. Full export generation (Markdown report, CSV table, JSON audit trail)
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import docx

from checksense import (
    HeuristicFusionScorer,
    Phase4Pipeline,
    apply_refinement_to_analysis_result,
    commit_approved_refinement,
    extract_from_docx,
    extract_from_txt,
    generate_audit_trail_json,
    generate_csv_srs_report,
    generate_markdown_srs_report,
    summarize_srs_analysis,
)
from checksense.document import analyze_srs_requirements


def test_end_to_end_srs_document_analysis_and_export_smoke_test(tmp_path: Path):
    """Perform verified end-to-end smoke test on a multi-requirement SRS document."""
    # 1. Synthesize an authentic SRS Word document
    doc = docx.Document()
    doc.add_heading("1.0 Scope & Purpose", level=1)
    doc.add_paragraph("This document specifies the software requirements for the Mission Control System.")

    doc.add_heading("2.0 Telemetry Processing Requirements", level=1)
    doc.add_paragraph("REQ-TEL-01: The telemetry module shall process telemetry frames quickly upon reception.")
    doc.add_paragraph("REQ-TEL-02: The system shall encrypt all outgoing command frames.")

    doc.add_heading("3.0 Thermal Subsystem", level=1)
    doc.add_paragraph("REQ-TH-01: The thermal controller shall maintain temperature between 18 and 24 degrees Celsius.")

    doc_bytes = io.BytesIO()
    doc.save(doc_bytes)
    content = doc_bytes.getvalue()

    # 2. Extract candidate requirements
    reqs, warnings = extract_from_docx(content, filename="MissionControl_SRS.docx")
    assert len(reqs) == 3
    assert len(warnings) == 0

    ids = [r.req_id for r in reqs]
    assert "REQ-TEL-01" in ids
    assert "REQ-TEL-02" in ids
    assert "REQ-TH-01" in ids

    # 3. Batch Defect Analysis & Fusion
    pipeline = Phase4Pipeline()
    scorer = HeuristicFusionScorer()

    results = analyze_srs_requirements(
        requirements=reqs,
        pipeline=pipeline,
        fusion_scorer=scorer,
        llm_client=None,  # Offline deterministic fallback
    )
    assert len(results) == 3

    # Verify REQ-TEL-01 has detected defect (vague term 'quickly')
    tel_01 = next(r for r in results if r.requirement.req_id == "REQ-TEL-01")
    assert tel_01.has_defects is True
    assert len(tel_01.questions) >= 1
    assert tel_01.clarity_heuristic < 1.00

    # 4. Formulate question & Stakeholder Clarification
    target_q = tel_01.questions[0]
    assert target_q.question.endswith("?")
    assert target_q.status == "open"

    # 5. Propose Refinement
    proposed_text, is_num_valid, audit_msg, method = apply_refinement_to_analysis_result(
        result=tel_01,
        question_id=target_q.question_id,
        stakeholder_answer="within 250 milliseconds",
        llm_client=None,
    )
    assert is_num_valid is True
    assert "within 250 milliseconds" in proposed_text
    assert tel_01.is_approved is False  # Must remain unapproved before explicit commit

    # 6. Explicit Stakeholder Approval Gate
    commit_approved_refinement(
        result=tel_01,
        question_id=target_q.question_id,
        stakeholder_answer="within 250 milliseconds",
        refined_text=proposed_text,
        method_used=method,
    )
    assert tel_01.is_approved is True
    assert tel_01.current_text == proposed_text
    assert target_q.status == "answered"
    assert len(tel_01.refinement_history) == 1

    # 7. Document-Level Aggregation
    summary = summarize_srs_analysis(results, filename="MissionControl_SRS.docx")
    assert summary.total_candidates == 3
    assert summary.total_analyzed == 3
    assert summary.approved_refinements_count == 1
    assert summary.macro_avg_iso_score > 0.0
    assert summary.macro_avg_clarity_score > 0.0

    # 8. Multi-format Report Generation & Verification
    md_report = generate_markdown_srs_report(summary, results)
    assert "# CheckSense SRS Document Analysis Report" in md_report
    assert "REQ-TEL-01" in md_report
    assert "within 250 milliseconds" in md_report
    assert "Approved by Stakeholder" in md_report
    assert "Traceability Matrix" in md_report

    csv_report = generate_csv_srs_report(results)
    assert "req_id,source_file,section" in csv_report
    assert "REQ-TEL-01,MissionControl_SRS.docx" in csv_report

    json_audit = generate_audit_trail_json(summary, results)
    parsed = json.loads(json_audit)
    assert parsed["summary"]["total_candidates"] == 3
    assert len(parsed["results"]) == 3
    assert parsed["results"][0]["is_approved"] is True


def test_txt_srs_document_smoke_test():
    """Verify TXT document batch analysis and export."""
    txt_content = """
# Section 1.0 General Specifications
REQ-SEC-01: The system shall authenticate user credentials using multi-factor authentication.
REQ-SEC-02: The system shall delete sensitive records immediately upon user request.
"""
    reqs, warnings = extract_from_txt(txt_content, filename="security_spec.txt")
    assert len(reqs) == 2
    assert len(warnings) == 0

    pipeline = Phase4Pipeline()
    scorer = HeuristicFusionScorer()
    results = analyze_srs_requirements(reqs, pipeline, scorer, llm_client=None)

    assert len(results) == 2
    summary = summarize_srs_analysis(results, filename="security_spec.txt")
    assert summary.total_analyzed == 2

    md = generate_markdown_srs_report(summary, results)
    assert "REQ-SEC-01" in md
    assert "REQ-SEC-02" in md

